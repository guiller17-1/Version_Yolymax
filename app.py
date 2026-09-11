import os
import re
import io
import json
import random
import string
from datetime import datetime, date
import requests
from openpyxl import Workbook
from flask import Flask, render_template, request, jsonify, redirect, url_for, session, send_file
from models import db, Order, OrderLine, StockSnapshot

app = Flask(__name__)

# Configuración de variables de entorno
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'clave_secreta_desarrollo')
ADMIN_PASSWORD = os.environ.get('ADMIN_PASSWORD', 'admin123')
CRON_SECRET = os.environ.get('CRON_SECRET', 'cron_secret_123')
ONEDRIVE_URL = os.environ.get('ONEDRIVE_URL', '')

# Configuración de Neon / PostgreSQL
db_url = os.environ.get('DATABASE_URL', 'sqlite:///app.db')
if db_url.startswith("postgres://"):
    db_url = db_url.replace("postgres://", "postgresql://", 1)
app.config['SQLALCHEMY_DATABASE_URI'] = db_url
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db.init_app(app)

with app.app_context():
    db.create_all()

# --- FUNCIONES AUXILIARES ---

def load_local_inventory():
    """Carga el inventario local desde inventory.json"""
    filepath = os.path.join(app.root_path, 'inventory.json')
    if os.path.exists(filepath):
        with open(filepath, 'r', encoding='utf-8') as f:
            return json.load(f)
    return []

def get_latest_stock_snapshots():
    """Obtiene la última captura de stock por producto"""
    latest_date = db.session.query(db.func.max(StockSnapshot.snapshot_date)).scalar()
    if not latest_date:
        return {}, None
    
    snapshots = StockSnapshot.query.filter_by(snapshot_date=latest_date).all()
    stock_map = {s.product_id: s.stock for s in snapshots}
    return stock_map, latest_date

def calculate_product_stock():
    """Calcula el stock disponible considerando las reservas activas (Pendiente y Confirmado)"""
    products = load_local_inventory()
    stock_map, latest_date = get_latest_stock_snapshots()
    
    # Calcular reservas activas
    reserved_query = db.session.query(
        OrderLine.product_id,
        db.func.sum(OrderLine.quantity).label('total_reserved')
    ).join(Order).filter(Order.status.in_(['Pendiente', 'Confirmado']))
    
    if latest_date:
        # Solo descontar pedidos posteriores o iguales a la captura
        reserved_query = reserved_query.filter(db.func.date(Order.created_at) >= latest_date)
        
    reserved_map = dict(reserved_query.group_by(OrderLine.product_id).all())
    
    result = []
    for p in products:
        p_id = str(p['id'])
        base_stock = stock_map.get(p_id, p.get('stock', 0))
        reserved = reserved_map.get(p_id, 0)
        available = max(0, base_stock - reserved)
        
        result.append({
            "id": p_id,
            "nombre": p.get('nombre', ''),
            "descripcion": p.get('descripcion', ''),
            "precio": float(p.get('precio', 0.0)),
            "imagen": p.get('imagen', ''),
            "stock_base": base_stock,
            "reservado": reserved,
            "stock_disponible": available
        })
    return result

def generate_order_code():
    """Genera código de pedido único: PED-YYYYMMDD-XXXXXX"""
    date_str = datetime.now().strftime('%Y%m%d')
    random_str = ''.join(random.choices(string.ascii_uppercase + string.digits, k=6))
    return f"PED-{date_str}-{random_str}"

# --- RUTAS PRINCIPALES DE LA APLICACIÓN ---

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/api/productos', methods=['GET'])
def api_productos():
    productos = calculate_product_stock()
    return jsonify(productos)

@app.route('/api/verificar-disponibilidad', methods=['POST'])
def api_verificar_disponibilidad():
    data = request.json or []
    productos = {p['id']: p['stock_disponible'] for p in calculate_product_stock()}
    
    detalles = []
    todo_disponible = True
    
    for item in data:
        p_id = str(item.get('product_id'))
        req_qty = int(item.get('quantity', 0))
        avail = productos.get(p_id, 0)
        
        es_valido = req_qty <= avail
        if not es_valido:
            todo_disponible = False
            
        detalles.append({
            "product_id": p_id,
            "solicitado": req_qty,
            "disponible": avail,
            "suficiente": es_valido
        })
        
    return jsonify({"disponible": todo_disponible, "detalles": detalles})

@app.route('/api/pedidos', methods=['POST'])
def api_crear_pedido():
    data = request.json or {}
    nombre = data.get('nombre', '').strip()
    telefono = data.get('telefono', '').strip()
    items = data.get('items', [])
    
    # Validaciones
    if len(nombre.split()) < 2:
        return jsonify({"error": "Debe ingresar su nombre completo (nombre y apellido)."}), 400
        
    if not re.match(r'^9\d{8}$', telefono):
        return jsonify({"error": "Debe ingresar un número de celular peruano válido (9 dígitos iniciando con 9)."}), 400
        
    if not items:
        return jsonify({"error": "El pedido no contiene productos."}), 400
        
    # Verificar stock
    productos_stock = {p['id']: p for p in calculate_product_stock()}
    total = 0.0
    lines_to_create = []
    
    for item in items:
        p_id = str(item.get('product_id'))
        qty = int(item.get('quantity', 0))
        
        if p_id not in productos_stock:
            return jsonify({"error": f"Producto ID {p_id} no existe."}), 400
            
        prod = productos_stock[p_id]
        if qty > prod['stock_disponible']:
            return jsonify({"error": f"Stock insuficiente para '{prod['nombre']}'. Disponible: {prod['stock_disponible']}"}), 400
            
        subtotal = prod['precio'] * qty
        total += subtotal
        lines_to_create.append({
            "product_id": p_id,
            "product_name": prod['nombre'],
            "unit_price": prod['precio'],
            "quantity": qty
        })
        
    # Crear Pedido
    codigo = generate_order_code()
    order = Order(
        code=codigo,
        customer_name=nombre,
        customer_phone=telefono,
        status='Pendiente',
        total_amount=total
    )
    db.session.add(order)
    db.session.flush()
    
    for l in lines_to_create:
        line = OrderLine(
            order_id=order.id,
            product_id=l['product_id'],
            product_name=l['product_name'],
            unit_price=l['unit_price'],
            quantity=l['quantity']
        )
        db.session.add(line)
        
    db.session.commit()
    return jsonify({"mensaje": "Pedido registrado exitosamente", "codigo": codigo}), 201

# --- TAREA PROGRAMADA (CRON) ---

@app.route('/api/tareas/captura-diaria', methods=['GET', 'POST'])
def captura_diaria():
    token = request.headers.get('X-Cron-Secret') or request.args.get('secret')
    if token != CRON_SECRET:
        return jsonify({"error": "No autorizado"}), 401
        
    today = date.today()
    productos_capturados = []
    
    # Intentar descargar Excel desde OneDrive
    if ONEDRIVE_URL:
        try:
            resp = requests.get(ONEDRIVE_URL, timeout=15)
            if resp.status_code == 200:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True)
                sheet = wb['Inventario'] if 'Inventario' in wb.sheetnames else wb.active
                
                # Eliminar registros previos de hoy si existen
                StockSnapshot.query.filter_by(snapshot_date=today).delete()
                
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if row and len(row) >= 3 and row[0]:
                        p_id = str(row[0]).strip()
                        p_name = str(row[1]).strip() if row[1] else p_id
                        stock_val = int(row[2]) if row[2] is not None else 0
                        
                        snap = StockSnapshot(
                            snapshot_date=today,
                            product_id=p_id,
                            product_name=p_name,
                            stock=stock_val
                        )
                        db.session.add(snap)
                        productos_capturados.append({"id": p_id, "stock": stock_val})
                
                db.session.commit()
                return jsonify({"status": "exito", "origen": "onedrive", "capturados": len(productos_capturados)})
        except Exception as e:
            app.logger.error(f"Error descargando OneDrive: {e}")
            
    # Fallback: Usar inventory.json
    local_inventory = load_local_inventory()
    StockSnapshot.query.filter_by(snapshot_date=today).delete()
    
    for item in local_inventory:
        snap = StockSnapshot(
            snapshot_date=today,
            product_id=str(item['id']),
            product_name=item.get('nombre', ''),
            stock=int(item.get('stock', 0))
        )
        db.session.add(snap)
        productos_capturados.append({"id": item['id'], "stock": item.get('stock', 0)})
        
    db.session.commit()
    return jsonify({"status": "exito", "origen": "fallback_json", "capturados": len(productos_capturados)})

# --- PANEL DE ADMINISTRACIÓN ---

@app.route('/admin/login', methods=['GET', 'POST'])
def admin_login():
    if request.method == 'POST':
        password = request.form.get('password')
        if password == ADMIN_PASSWORD:
            session['admin_logged_in'] = True
            return redirect(url_for('admin_orders'))
        return render_template('admin_login.html', error="Contraseña incorrecta")
    return render_template('admin_login.html')

@app.route('/admin/logout')
def admin_logout():
    session.pop('admin_logged_in', None)
    return redirect(url_for('admin_login'))

@app.route('/admin/orders')
def admin_orders():
    if not session.get('admin_logged_in'):
        return redirect(url_for('admin_login'))
        
    status_filter = request.args.get('status')
    query = Order.query
    if status_filter in ['Pendiente', 'Confirmado', 'Rechazado']:
        query = query.filter_by(status=status_filter)
        
    orders = query.order_by(Order.created_at.desc()).all()
    return render_template('admin_orders.html', orders=orders, current_status=status_filter)

@app.route('/admin/orders/<int:order_id>/status', methods=['POST'])
def update_order_status(order_id):
    if not session.get('admin_logged_in'):
        return jsonify({"error": "No autorizado"}), 401
        
    order = Order.query.get_or_404(order_id)
    new_status = request.form.get('status')
    if new_status in ['Pendiente', 'Confirmado', 'Rechazado']:
        order.status = new_status
        db.session.commit()
        
    return redirect(url_for('admin_orders'))

@app.route('/admin/exportar-pedidos.xlsx')
def export_orders_excel():
    if not session.get('admin_logged_in'):
        return redirect(url_for('admin_login'))
        
    wb = Workbook()
    
    # 1. Pestaña Pedidos
    ws_orders = wb.active
    ws_orders.title = "Pedidos"
    ws_orders.append(["ID", "Código", "Cliente", "Teléfono", "Estado", "Total (S/)", "Fecha Registro"])
    
    orders = Order.query.order_by(Order.created_at.desc()).all()
    for o in orders:
        ws_orders.append([o.id, o.code, o.customer_name, o.customer_phone, o.status, o.total_amount, o.created_at.strftime('%Y-%m-%d %H:%M:%S')])
        
    # 2. Pestaña DetallePedidos
    ws_details = wb.create_sheet(title="DetallePedidos")
    ws_details.append(["Código Pedido", "ID Producto", "Nombre Producto", "Precio Unitario (S/)", "Cantidad", "Subtotal (S/)"])
    
    for o in orders:
        for line in o.lines:
            subtotal = line.unit_price * line.quantity
            ws_details.append([o.code, line.product_id, line.product_name, line.unit_price, line.quantity, subtotal])
            
    # 3. Pestaña StockDiario
    ws_stock = wb.create_sheet(title="StockDiario")
    ws_stock.append(["Fecha Captura", "ID Producto", "Nombre Producto", "Stock Inicial", "Reservado", "Stock Disponible"])
    
    current_stocks = calculate_product_stock()
    latest_date = db.session.query(db.func.max(StockSnapshot.snapshot_date)).scalar() or date.today()
    for s in current_stocks:
        ws_stock.append([latest_date.strftime('%Y-%m-%d'), s['id'], s['nombre'], s['stock_base'], s['reservado'], s['stock_disponible']])
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"Reporte_Pedidos_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(output, download_name=filename, as_attachment=True)

if __name__ == '__main__':
    app.run(debug=True)
