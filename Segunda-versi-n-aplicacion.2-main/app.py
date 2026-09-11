import io
import json
import os
import re
import secrets
from datetime import datetime, time as dt_time, timezone
from decimal import Decimal
from functools import wraps
from pathlib import Path
from zoneinfo import ZoneInfo

import requests
from flask import (
    Flask,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import Workbook, load_workbook
from sqlalchemy import func

from models import Order, OrderLine, StockSnapshot, db


# =========================================================
# ZONAS HORARIAS
# =========================================================

LIMA = ZoneInfo("America/Lima")
UTC = timezone.utc


# =========================================================
# CONFIGURACIÓN DE FLASK
# =========================================================

app = Flask(__name__)

app.config["SECRET_KEY"] = os.getenv(
    "SECRET_KEY",
    secrets.token_hex(32)
)

database_url = os.getenv(
    "DATABASE_URL",
    "sqlite:///pedidos.db"
)

if database_url.startswith("postgres://"):
    database_url = database_url.replace(
        "postgres://",
        "postgresql://",
        1
    )

app.config["SQLALCHEMY_DATABASE_URI"] = database_url
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

db.init_app(app)


ONEDRIVE_URL = os.getenv(
    "ONEDRIVE_URL",
    ""
).strip()

ADMIN_PASSWORD = os.getenv(
    "ADMIN_PASSWORD",
    ""
).strip()

CRON_SECRET = os.getenv(
    "CRON_SECRET",
    ""
).strip()


# =========================================================
# UTILIDADES DE FECHA Y HORA
# =========================================================

def ahora_utc():
    """
    Devuelve la fecha y hora actual en UTC
    incluyendo la información de zona horaria.
    """
    return datetime.now(UTC)


def ahora_lima():
    """
    Devuelve la fecha y hora actual de Perú.
    """
    return datetime.now(LIMA)


def asegurar_utc(fecha):
    """
    Convierte una fecha a UTC.

    Las fechas antiguas que no tienen zona horaria
    se interpretan como UTC.
    """
    if fecha is None:
        return None

    if fecha.tzinfo is None:
        return fecha.replace(tzinfo=UTC)

    return fecha.astimezone(UTC)


def convertir_a_hora_peru(fecha):
    """
    Convierte una fecha UTC a la hora de Perú.
    """
    fecha_utc = asegurar_utc(fecha)

    if fecha_utc is None:
        return None

    return fecha_utc.astimezone(LIMA)


@app.template_filter("hora_peru")
def filtro_hora_peru(fecha):
    """
    Ejemplo en HTML:

    {{ order.created_at | hora_peru }}
    """
    fecha_peru = convertir_a_hora_peru(fecha)

    if fecha_peru is None:
        return ""

    return fecha_peru.strftime(
        "%d/%m/%Y %I:%M:%S %p"
    )


@app.template_filter("fecha_peru")
def filtro_fecha_peru(fecha):
    fecha_peru = convertir_a_hora_peru(fecha)

    if fecha_peru is None:
        return ""

    return fecha_peru.strftime("%d/%m/%Y")


@app.template_filter("hora_peru_corta")
def filtro_hora_peru_corta(fecha):
    fecha_peru = convertir_a_hora_peru(fecha)

    if fecha_peru is None:
        return ""

    return fecha_peru.strftime("%I:%M:%S %p")


# =========================================================
# UTILIDADES GENERALES
# =========================================================

def normalizar(texto):
    return str(texto or "").strip().casefold()


def entero_stock(valor):
    try:
        return max(
            0,
            int(float(valor or 0))
        )
    except (TypeError, ValueError):
        return 0


def estado(stock):
    if stock >= 3:
        return "Disponible"

    if stock >= 1:
        return "Ultimas unidades"

    return "No disponible"


# =========================================================
# DESCARGA DEL EXCEL DE ONEDRIVE
# =========================================================

def download_onedrive():
    if not ONEDRIVE_URL:
        raise RuntimeError(
            "ONEDRIVE_URL no configurada"
        )

    candidates = [ONEDRIVE_URL]

    separator = (
        "&"
        if "?" in ONEDRIVE_URL
        else "?"
    )

    candidates.append(
        ONEDRIVE_URL + separator + "download=1"
    )

    last_error = None

    for url in candidates:
        try:
            response = requests.get(
                url,
                timeout=30,
                allow_redirects=True
            )

            response.raise_for_status()

            if response.content[:2] == b"PK":
                return response.content

            last_error = ValueError(
                "OneDrive devolvió una página "
                "y no el archivo Excel"
            )

        except Exception as error:
            last_error = error

    raise last_error or RuntimeError(
        "No se pudo descargar el Excel"
    )


# =========================================================
# LECTURA DEL EXCEL
# =========================================================

def parse_excel(content):
    workbook = load_workbook(
        io.BytesIO(content),
        data_only=True,
        read_only=True,
        keep_vba=True
    )

    if "Inventario" not in workbook.sheetnames:
        raise ValueError(
            "No existe la hoja Inventario"
        )

    worksheet = workbook["Inventario"]

    rows = worksheet.iter_rows(
        values_only=True
    )

    first = next(rows, None)

    if first is None:
        raise ValueError(
            "La hoja Inventario está vacía"
        )

    headers = {
        str(value).strip(): index
        for index, value in enumerate(first)
        if value is not None
    }

    required = (
        "Marca",
        "Producto",
        "Stock Actual",
        "Precio",
        "Imagen 1",
        "Imagen 2",
        "Imagen 3"
    )

    for column in required:
        if column not in headers:
            raise ValueError(
                f"No existe la columna {column}"
            )

    items = []

    for row in rows:
        marca = row[headers["Marca"]]
        producto = row[headers["Producto"]]

        if not marca or not producto:
            continue

        imagenes = []

        for column in (
            "Imagen 1",
            "Imagen 2",
            "Imagen 3"
        ):
            value = str(
                row[headers[column]] or ""
            ).strip()

            if value:
                imagenes.append(value)

        try:
            precio = Decimal(
                str(
                    row[headers["Precio"]] or 0
                )
            )
        except Exception:
            precio = Decimal("0")

        items.append({
            "marca": str(marca).strip(),
            "producto": str(producto).strip(),
            "precio": precio,
            "stock": entero_stock(
                row[headers["Stock Actual"]]
            ),
            "imagenes": imagenes,
        })

    return items


# =========================================================
# INVENTARIO DE RESPALDO
# =========================================================

def load_fallback():
    path = Path(
        app.root_path,
        "inventory.json"
    )

    if not path.exists():
        return []

    raw = json.loads(
        path.read_text(
            encoding="utf-8"
        )
    )

    result = []

    for item in raw:
        imagenes = item.get(
            "imagenes",
            []
        )

        if not isinstance(imagenes, list):
            imagenes = []

        try:
            precio = Decimal(
                str(
                    item.get("precio", 0) or 0
                )
            )
        except Exception:
            precio = Decimal("0")

        result.append({
            "marca": str(
                item.get("marca", "")
            ).strip(),
            "producto": str(
                item.get("producto", "")
            ).strip(),
            "precio": precio,
            "stock": entero_stock(
                item.get("stock", 0)
            ),
            "imagenes": imagenes,
        })

    return result


# =========================================================
# CAPTURAS DIARIAS DEL STOCK
# =========================================================

def snapshot_date_in_use():
    today = ahora_lima().date()

    today_exists = (
        db.session.query(StockSnapshot.id)
        .filter_by(snapshot_date=today)
        .first()
    )

    if today_exists:
        return today

    return db.session.query(
        func.max(
            StockSnapshot.snapshot_date
        )
    ).scalar()


def create_daily_snapshot(force=False):
    today = ahora_lima().date()

    exists = (
        db.session.query(StockSnapshot.id)
        .filter_by(snapshot_date=today)
        .first()
    )

    if exists and not force:
        return {
            "created": False,
            "date": today.isoformat(),
            "message": (
                "La captura de hoy ya existe"
            )
        }

    try:
        items = parse_excel(
            download_onedrive()
        )

        source = "OneDrive"

    except Exception as error:
        app.logger.exception(
            "Fallo al descargar el inventario: %s",
            error
        )

        if exists:
            raise

        items = load_fallback()
        source = "respaldo"

    if not items:
        raise RuntimeError(
            "No hay productos para crear "
            "la captura diaria"
        )

    if force:
        StockSnapshot.query.filter_by(
            snapshot_date=today
        ).delete(
            synchronize_session=False
        )

    captured_at_utc = ahora_utc()

    for item in items:
        db.session.add(
            StockSnapshot(
                snapshot_date=today,
                captured_at=captured_at_utc,
                marca=item["marca"],
                producto=item["producto"],
                precio=item["precio"],
                stock=item["stock"],
                imagenes_json=json.dumps(
                    item["imagenes"],
                    ensure_ascii=False
                ),
            )
        )

    db.session.commit()

    return {
        "created": True,
        "date": today.isoformat(),
        "hora_peru": ahora_lima().strftime(
            "%d/%m/%Y %I:%M:%S %p"
        ),
        "source": source,
        "products": len(items)
    }


def ensure_snapshot():
    current_date = snapshot_date_in_use()
    now_peru = ahora_lima()

    if current_date is None:
        return create_daily_snapshot(
            force=False
        )

    if (
        current_date < now_peru.date()
        and now_peru.time() >= dt_time(6, 0)
    ):
        return create_daily_snapshot(
            force=False
        )

    return None


# =========================================================
# RESERVAS Y STOCK DISPONIBLE
# =========================================================

def obtener_primera_captura(date_used):
    if date_used is None:
        return None

    return (
        db.session.query(
            func.min(
                StockSnapshot.captured_at
            )
        )
        .filter(
            StockSnapshot.snapshot_date
            == date_used
        )
        .scalar()
    )


def reserved_quantities():
    date_used = snapshot_date_in_use()

    if date_used is None:
        return {}

    first_capture = obtener_primera_captura(
        date_used
    )

    query = (
        db.session.query(
            OrderLine.marca,
            OrderLine.producto,
            func.sum(OrderLine.quantity)
        )
        .join(Order)
    )

    if first_capture is not None:
        query = query.filter(
            (Order.status == "Pendiente")
            |
            (
                (Order.status == "Confirmado")
                &
                (
                    Order.created_at
                    >= first_capture
                )
            )
        )
    else:
        query = query.filter(
            Order.status == "Pendiente"
        )

    rows = (
        query
        .group_by(
            OrderLine.marca,
            OrderLine.producto
        )
        .all()
    )

    return {
        (
            normalizar(marca),
            normalizar(producto)
        ): int(quantity or 0)
        for marca, producto, quantity in rows
    }


def catalog_items():
    ensure_snapshot()

    date_used = snapshot_date_in_use()

    if date_used is None:
        return [], None

    reservations = reserved_quantities()

    snapshots = (
        StockSnapshot.query
        .filter_by(snapshot_date=date_used)
        .order_by(
            StockSnapshot.marca,
            StockSnapshot.producto
        )
        .all()
    )

    items = []

    for item in snapshots:
        reserved = reservations.get(
            (
                normalizar(item.marca),
                normalizar(item.producto)
            ),
            0
        )

        available = max(
            0,
            item.stock - reserved
        )

        try:
            imagenes = json.loads(
                item.imagenes_json or "[]"
            )
        except (TypeError, json.JSONDecodeError):
            imagenes = []

        items.append({
            "marca": item.marca,
            "producto": item.producto,
            "precio": float(item.precio),
            "imagenes": imagenes,
            "estado": estado(available),
        })

    return items, date_used


def find_snapshot(marca, producto):
    date_used = snapshot_date_in_use()

    if date_used is None:
        return None

    return StockSnapshot.query.filter(
        StockSnapshot.snapshot_date
        == date_used,
        func.lower(
            StockSnapshot.marca
        ) == str(marca or "").strip().lower(),
        func.lower(
            StockSnapshot.producto
        ) == str(producto or "").strip().lower(),
    ).first()


def available_stock(snapshot):
    first_capture = obtener_primera_captura(
        snapshot.snapshot_date
    )

    query = (
        db.session.query(
            func.coalesce(
                func.sum(OrderLine.quantity),
                0
            )
        )
        .join(Order)
        .filter(
            func.lower(
                OrderLine.marca
            ) == snapshot.marca.lower(),
            func.lower(
                OrderLine.producto
            ) == snapshot.producto.lower(),
        )
    )

    if first_capture is not None:
        query = query.filter(
            (Order.status == "Pendiente")
            |
            (
                (Order.status == "Confirmado")
                &
                (
                    Order.created_at
                    >= first_capture
                )
            )
        )
    else:
        query = query.filter(
            Order.status == "Pendiente"
        )

    reserved = query.scalar()

    return max(
        0,
        snapshot.stock - int(reserved or 0)
    )


# =========================================================
# PROTECCIÓN DEL PANEL ADMINISTRATIVO
# =========================================================

def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("admin"):
            return redirect(
                url_for(
                    "admin_login",
                    next=request.path
                )
            )

        return view(*args, **kwargs)

    return wrapped


# =========================================================
# INICIALIZACIÓN DE LA BASE
# =========================================================

@app.before_request
def initialize_database():
    db.create_all()


# =========================================================
# CATÁLOGO
# =========================================================

@app.get("/")
def home():
    return render_template("index.html")


@app.get("/api/productos")
def productos():
    items, date_used = catalog_items()

    return jsonify({
        "productos": items,
        "fecha_stock": (
            date_used.isoformat()
            if date_used
            else None
        ),
        "hora_peru": ahora_lima().strftime(
            "%d/%m/%Y %I:%M:%S %p"
        ),
    })


# =========================================================
# VERIFICACIÓN DE DISPONIBILIDAD
# =========================================================

@app.post("/api/verificar-disponibilidad")
def verificar_disponibilidad():
    data = request.get_json(
        silent=True
    ) or {}

    try:
        quantity = int(
            data.get("cantidad", 0)
        )

        in_cart = int(
            data.get(
                "cantidad_en_cesta",
                0
            )
        )

    except (TypeError, ValueError):
        quantity = 0
        in_cart = 0

    if quantity < 1 or quantity > 99:
        return jsonify({
            "disponible": False,
            "mensaje": (
                "Selecciona una cantidad válida."
            )
        }), 400

    ensure_snapshot()

    snapshot = find_snapshot(
        data.get("marca"),
        data.get("producto")
    )

    if snapshot is None:
        return jsonify({
            "disponible": False,
            "mensaje": "Producto no encontrado."
        }), 404

    requested_total = (
        quantity + max(0, in_cart)
    )

    if requested_total > available_stock(snapshot):
        return jsonify({
            "disponible": False,
            "mensaje": (
                "No hay disponibilidad suficiente. "
                "Prueba con una cantidad menor."
            ),
        })

    return jsonify({
        "disponible": True,
        "mensaje": "Cantidad disponible."
    })


# =========================================================
# CREACIÓN DE PEDIDOS
# =========================================================

@app.post("/api/pedidos")
def crear_pedido():
    payload = request.get_json(
        silent=True
    ) or {}

    name = str(
        payload.get("nombre", "")
    ).strip()

    phone = re.sub(
        r"\D",
        "",
        str(payload.get("telefono", ""))
    )

    lines = payload.get(
        "productos",
        []
    )

    if len(name) < 3:
        return jsonify({
            "ok": False,
            "mensaje": (
                "Ingresa tu nombre completo."
            )
        }), 400

    if not re.fullmatch(r"9\d{8}", phone):
        return jsonify({
            "ok": False,
            "mensaje": (
                "Ingresa un celular peruano "
                "válido de 9 dígitos."
            )
        }), 400

    if not isinstance(lines, list) or not lines:
        return jsonify({
            "ok": False,
            "mensaje": "La cesta está vacía."
        }), 400

    ensure_snapshot()

    prepared = []
    total = Decimal("0")

    try:
        for line in lines:
            quantity = int(
                line.get("cantidad", 0)
            )

            if quantity < 1 or quantity > 99:
                raise ValueError(
                    "Cantidad inválida"
                )

            snapshot = find_snapshot(
                line.get("marca"),
                line.get("producto")
            )

            if (
                snapshot is None
                or quantity
                > available_stock(snapshot)
            ):
                return jsonify({
                    "ok": False,
                    "mensaje": (
                        "Uno de los productos ya no "
                        "tiene disponibilidad suficiente."
                    ),
                }), 409

            subtotal = (
                snapshot.precio * quantity
            )

            total += subtotal

            prepared.append(
                (
                    snapshot,
                    quantity,
                    subtotal
                )
            )

        code = (
            f"PED-"
            f"{ahora_lima().strftime('%Y%m%d')}-"
            f"{secrets.token_hex(3).upper()}"
        )

        current_utc = ahora_utc()

        order = Order(
            code=code,
            customer_name=name,
            customer_phone=phone,
            total=total,
            status="Pendiente",
            channel="WhatsApp",
            created_at=current_utc,
            updated_at=current_utc,
        )

        db.session.add(order)
        db.session.flush()

        for snapshot, quantity, subtotal in prepared:
            db.session.add(
                OrderLine(
                    order_id=order.id,
                    marca=snapshot.marca,
                    producto=snapshot.producto,
                    quantity=quantity,
                    unit_price=snapshot.precio,
                    subtotal=subtotal,
                )
            )

        db.session.commit()

    except Exception as error:
        db.session.rollback()

        app.logger.exception(
            "No se pudo crear el pedido: %s",
            error
        )

        return jsonify({
            "ok": False,
            "mensaje": (
                "No se pudo registrar el pedido."
            )
        }), 500

    return jsonify({
        "ok": True,
        "codigo": order.code,
        "estado": order.status,
        "total": float(order.total),
        "fecha_peru": filtro_hora_peru(
            order.created_at
        ),
    }), 201


# =========================================================
# ACCESO ADMINISTRATIVO
# =========================================================

@app.get("/admin/login")
def admin_login():
    return render_template(
        "admin_login.html",
        error=None
    )


@app.post("/admin/login")
def admin_login_post():
    password = request.form.get(
        "password",
        ""
    )

    if (
        not ADMIN_PASSWORD
        or not secrets.compare_digest(
            password,
            ADMIN_PASSWORD
        )
    ):
        return render_template(
            "admin_login.html",
            error="Contraseña incorrecta"
        ), 401

    session.clear()
    session["admin"] = True

    return redirect(
        url_for("admin_orders")
    )


@app.post("/admin/logout")
def admin_logout():
    session.clear()

    return redirect(
        url_for("admin_login")
    )


# =========================================================
# PANEL DE PEDIDOS
# =========================================================

@app.get("/admin/pedidos")
@admin_required
def admin_orders():
    orders = (
        Order.query
        .order_by(
            Order.created_at.desc()
        )
        .all()
    )

    return render_template(
        "admin_orders.html",
        orders=orders
    )


@app.post(
    "/admin/pedidos/<int:order_id>/estado"
)
@admin_required
def update_order_status(order_id):
    order = db.session.get(
        Order,
        order_id
    )

    if order is None:
        return jsonify({
            "ok": False,
            "mensaje": "Pedido no encontrado"
        }), 404

    new_status = str(
        (
            request.get_json(
                silent=True
            ) or {}
        ).get("estado", "")
    ).strip()

    valid_statuses = {
        "Pendiente",
        "Confirmado",
        "Rechazado"
    }

    if new_status not in valid_statuses:
        return jsonify({
            "ok": False,
            "mensaje": "Estado inválido"
        }), 400

    order.status = new_status
    order.updated_at = ahora_utc()

    db.session.commit()

    return jsonify({
        "ok": True,
        "estado": order.status,
        "actualizado_peru": filtro_hora_peru(
            order.updated_at
        ),
    })


# =========================================================
# EXPORTACIÓN A EXCEL
# =========================================================

@app.get("/admin/exportar-pedidos.xlsx")
@admin_required
def export_orders():
    workbook = Workbook()

    orders_sheet = workbook.active
    orders_sheet.title = "Pedidos"

    orders_sheet.append([
        "Código",
        "Fecha Perú",
        "Hora Perú",
        "Cliente",
        "Celular",
        "Total",
        "Estado",
        "Canal",
        "Última actualización Perú",
    ])

    details_sheet = workbook.create_sheet(
        "DetallePedidos"
    )

    details_sheet.append([
        "Código",
        "Fecha Perú",
        "Hora Perú",
        "Marca",
        "Producto",
        "Cantidad",
        "Precio unitario",
        "Subtotal",
        "Estado",
    ])

    orders = (
        Order.query
        .order_by(
            Order.created_at.asc()
        )
        .all()
    )

    for order in orders:
        created_peru = convertir_a_hora_peru(
            order.created_at
        )

        updated_peru = convertir_a_hora_peru(
            order.updated_at
        )

        orders_sheet.append([
            order.code,
            (
                created_peru.strftime("%d/%m/%Y")
                if created_peru
                else ""
            ),
            (
                created_peru.strftime("%I:%M:%S %p")
                if created_peru
                else ""
            ),
            order.customer_name,
            order.customer_phone,
            float(order.total),
            order.status,
            order.channel,
            (
                updated_peru.strftime(
                    "%d/%m/%Y %I:%M:%S %p"
                )
                if updated_peru
                else ""
            ),
        ])

        for line in order.lines:
            details_sheet.append([
                order.code,
                (
                    created_peru.strftime("%d/%m/%Y")
                    if created_peru
                    else ""
                ),
                (
                    created_peru.strftime("%I:%M:%S %p")
                    if created_peru
                    else ""
                ),
                line.marca,
                line.producto,
                line.quantity,
                float(line.unit_price),
                float(line.subtotal),
                order.status,
            ])

    snapshots_sheet = workbook.create_sheet(
        "StockDiario"
    )

    snapshots_sheet.append([
        "Fecha de stock",
        "Fecha de captura Perú",
        "Hora de captura Perú",
        "Marca",
        "Producto",
        "Precio",
        "Stock base",
    ])

    snapshots = (
        StockSnapshot.query
        .order_by(
            StockSnapshot.snapshot_date.asc(),
            StockSnapshot.marca.asc(),
            StockSnapshot.producto.asc(),
        )
        .all()
    )

    for snapshot in snapshots:
        captured_peru = convertir_a_hora_peru(
            snapshot.captured_at
        )

        snapshots_sheet.append([
            snapshot.snapshot_date.strftime(
                "%d/%m/%Y"
            ),
            (
                captured_peru.strftime("%d/%m/%Y")
                if captured_peru
                else ""
            ),
            (
                captured_peru.strftime("%I:%M:%S %p")
                if captured_peru
                else ""
            ),
            snapshot.marca,
            snapshot.producto,
            float(snapshot.precio),
            snapshot.stock,
        ])

    for worksheet in workbook.worksheets:
        for column_cells in worksheet.columns:
            max_length = 0

            column_letter = (
                column_cells[0].column_letter
            )

            for cell in column_cells:
                value_length = len(
                    str(cell.value or "")
                )

                if value_length > max_length:
                    max_length = value_length

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max_length + 2,
                45
            )

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

    output = io.BytesIO()

    workbook.save(output)
    output.seek(0)

    filename = (
        "pedidos_"
        f"{ahora_lima().strftime('%Y%m%d_%H%M')}"
        ".xlsx"
    )

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
    )


# =========================================================
# TAREA DIARIA DE LAS 6:00 A. M. DE PERÚ
# =========================================================

@app.post("/api/tareas/captura-diaria")
def daily_snapshot_task():
    supplied = request.headers.get(
        "X-Cron-Secret",
        ""
    )

    if (
        not CRON_SECRET
        or not secrets.compare_digest(
            supplied,
            CRON_SECRET
        )
    ):
        return jsonify({
            "ok": False
        }), 403

    try:
        result = create_daily_snapshot(
            force=False
        )

        return jsonify({
            "ok": True,
            "hora_peru": ahora_lima().strftime(
                "%d/%m/%Y %I:%M:%S %p"
            ),
            **result
        })

    except Exception as error:
        app.logger.exception(
            "Fallo de captura diaria: %s",
            error
        )

        return jsonify({
            "ok": False,
            "mensaje": str(error),
            "hora_peru": ahora_lima().strftime(
                "%d/%m/%Y %I:%M:%S %p"
            ),
        }), 500


# =========================================================
# ESTADO DEL SERVICIO
# =========================================================

@app.get("/health")
def health():
    return {
        "ok": True,
        "zona_horaria": "America/Lima",
        "fecha_hora_peru": ahora_lima().strftime(
            "%d/%m/%Y %I:%M:%S %p"
        ),
        "fecha_hora_utc": ahora_utc().isoformat(),
    }


# =========================================================
# CREACIÓN INICIAL DE TABLAS
# =========================================================

with app.app_context():
    db.create_all()


if __name__ == "__main__":
    app.run(
        host="0.0.0.0",
        port=int(
            os.getenv("PORT", "5000")
        ),
        debug=True
    )
